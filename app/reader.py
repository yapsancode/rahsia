"""Turn an uploaded file into plain text. It runs on this laptop and keeps nothing on disk.

The text goes back into the notes box, so the person sees it first. It then
takes the same path as typed notes, including the check for hidden instructions.
"""
import io

from docx import Document
from pypdf import PdfReader

MAX_BYTES = 5_000_000
MAX_CHARS = 20_000  # far more than the builder reads; it stops one big file from slowing the page
BOMS = (b"\xef\xbb\xbf", b"\xff\xfe", b"\xfe\xff")


class FileError(ValueError):
    """The message is shown to the person, so it says what to do."""


def _pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise FileError("This PDF is locked. Remove the password, then try again.")
    pages: list[str] = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
        if sum(map(len, pages)) >= MAX_CHARS:
            break
    return "\n".join(pages)


def _docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        lines += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    return "\n".join(lines)


def _plain(data: bytes) -> str:
    if data.startswith(BOMS[1:]):
        return data.decode("utf-16")
    if b"\0" in data[:8192]:
        raise FileError("This is not a text file. Try .txt, .md, .csv, .pdf or .docx.")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def read_file(kind: str, data: bytes) -> str:
    """`kind` is the file ending, such as 'pdf'. Anything but pdf and docx is read as text."""
    if not data:
        raise FileError("This file is empty.")
    if len(data) > MAX_BYTES:
        raise FileError("This file is over 5 MB. Copy the part you need into the notes box.")
    kind = kind.lower().lstrip(".")
    try:
        text = _pdf(data) if kind == "pdf" else _docx(data) if kind == "docx" else _plain(data)
    except FileError:
        raise
    except Exception as error:  # these libraries raise many types on a broken file
        raise FileError("This file could not be read. It may be damaged.") from error
    text = text.strip()
    if not text:
        raise FileError("No text found. A scanned PDF or a photo has no text to read.")
    return text[:MAX_CHARS]
