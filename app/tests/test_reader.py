import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.reader import MAX_BYTES, MAX_CHARS, FileError, read_file


def test_plain_text_of_any_ending():
    assert read_file("txt", "kami buat sensor".encode()) == "kami buat sensor"
    assert read_file(".MD", b"# notes") == "# notes"
    assert read_file("", b"no ending") == "no ending"


def test_windows_encodings():
    assert read_file("txt", "﻿notes".encode("utf-8")) == "notes"
    assert read_file("txt", "notes".encode("utf-16")) == "notes"
    assert read_file("txt", "caf\xe9".encode("cp1252")) == "caf\xe9"


def test_docx_with_a_table():
    document = Document()
    document.add_paragraph("We raise RM300k.")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Malaysia"
    data = io.BytesIO()
    document.save(data)
    text = read_file("docx", data.getvalue())
    assert "We raise RM300k." in text and "Malaysia" in text


def test_a_pdf_with_no_text_says_so():
    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    data = io.BytesIO()
    writer.write(data)
    with pytest.raises(FileError, match="No text found"):
        read_file("pdf", data.getvalue())


def test_a_locked_pdf_says_so():
    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    writer.encrypt("fake-test-password")
    data = io.BytesIO()
    writer.write(data)
    with pytest.raises(FileError, match="locked"):
        read_file("pdf", data.getvalue())


@pytest.mark.parametrize("kind, data, message", [
    ("txt", b"", "empty"),
    ("exe", b"MZ\0\0\0binary", "not a text file"),
    ("pdf", b"this is not a pdf", "could not be read"),
    ("docx", b"this is not a docx", "could not be read"),
    ("txt", b"x" * (MAX_BYTES + 1), "over 5 MB"),
])
def test_bad_files_give_a_clear_message(kind, data, message):
    with pytest.raises(FileError, match=message):
        read_file(kind, data)


def test_very_long_text_is_cut():
    assert len(read_file("txt", b"x" * (MAX_CHARS + 50))) == MAX_CHARS
